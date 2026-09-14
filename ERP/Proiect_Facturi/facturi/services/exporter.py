# -*- coding: utf-8 -*-
import csv

from openpyxl import Workbook

from ERP.Proiect_Facturi.facturi.db import repository
from ERP.Proiect_Facturi.facturi.models.invoice_table_model import HEADERS

EXPORT_COLUMNS = [c for c in repository.COLUMNS if c != "id"]
EXPORT_BATCH = 5000


def _iter_rows(filters):
    offset = 0
    while True:
        batch = repository.fetch_page(
            filters, order_by="data_factura", order_dir="DESC", offset=offset, limit=EXPORT_BATCH
        )
        if not batch:
            break
        for row in batch:
            yield [row.get(c) for c in EXPORT_COLUMNS]
        offset += len(batch)


def export_to_csv(path: str, filters: dict = None) -> int:
    count = 0
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([HEADERS[c] for c in EXPORT_COLUMNS])
        for row in _iter_rows(filters):
            writer.writerow(row)
            count += 1
    return count


def export_to_excel(path: str, filters: dict = None) -> int:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Facturi")
    ws.append([HEADERS[c] for c in EXPORT_COLUMNS])
    count = 0
    for row in _iter_rows(filters):
        ws.append(row)
        count += 1
    wb.save(path)
    return count


def export(path: str, filters: dict = None) -> int:
    if path.lower().endswith(".csv"):
        return export_to_csv(path, filters)
    return export_to_excel(path, filters)
